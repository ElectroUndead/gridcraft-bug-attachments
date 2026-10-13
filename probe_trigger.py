"""Which property of the real workbook makes gridcraft-cli sort A2:.. instead of A1:.. ?"""
import os, subprocess
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.formatting.rule import CellIsRule

CLI = r"C:\Program Files\GridCraft\gridcraft-cli.exe"
OUT = r"C:\Users\ElectroUndead\Documents\deepseek-harness\Main\gridcraft-bug\repro\_probe"
os.makedirs(OUT, exist_ok=True)


def build(name, n_data=6, n_cols=3, styled=False, cf=False, extra_sheet=False, rows_declared=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Payments"
    ws.sheet_view.tabSelected = True
    ws["A1"] = "Payments from the main balance 01.10.2026-13.10.2026: %d rows" % n_data
    hdr = ["Date", "Company", "Amount"] + ["C%d" % i for i in range(4, n_cols + 1)]
    for c, v in enumerate(hdr, start=1):
        cell = ws.cell(row=2, column=c, value=v)
        if styled:
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
            cell.font = Font(bold=True)
    for i in range(n_data):
        ws.cell(row=3 + i, column=1, value="2026-10-%02d" % ((i % 28) + 1))
        ws.cell(row=3 + i, column=2, value="Company %03d" % i)
        ws.cell(row=3 + i, column=3, value=100 + i)
    last = 2 + n_data
    col = chr(64 + n_cols)
    ws.auto_filter.ref = "A2:%s%d" % (col, last)
    ws.freeze_panes = "A3"
    if cf:
        ws.conditional_formatting.add("C3:C%d" % last,
                                      CellIsRule(operator="greaterThan", formula=["300"], fill=PatternFill("solid", fgColor="FFC7CE")))
    if extra_sheet:
        ws2 = wb.create_sheet("Other")
        ws2["A1"] = "x"
    p = os.path.join(OUT, name + ".xlsx")
    wb.save(p)
    return p


cases = [
    ("m_6x3_plain",     dict()),
    ("m_6x3_styled",    dict(styled=True)),
    ("m_6x3_cf",        dict(cf=True)),
    ("m_6x3_2sheets",   dict(extra_sheet=True)),
    ("m_200x3",         dict(n_data=200)),
    ("m_236x3",         dict(n_data=236)),
    ("m_236x21",        dict(n_data=236, n_cols=21)),
    ("m_236x21_styled", dict(n_data=236, n_cols=21, styled=True, cf=True, extra_sheet=True)),
]

for name, kw in cases:
    p = build(name, **kw)
    r = subprocess.run([CLI, "run", "--in", p, "--cmd", 'selection.set={"cell":"B2"}',
                        "--cmd", 'data.sortAscending={"column":"B"}',
                        "--out", os.path.join(OUT, name + "_out.xlsx")],
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    lines = [l for l in (r.stdout or "").strip().splitlines() if "sorted" in l]
    print(f"{name:18s} -> {lines[0] if lines else (r.stderr or '')[:70]}")
